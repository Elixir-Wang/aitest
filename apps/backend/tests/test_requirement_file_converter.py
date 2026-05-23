from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.requirement_file_converter import convert_requirement_file_to_markdown


class RequirementFileConverterTest(unittest.TestCase):
    def test_markdown_file_is_saved_directly(self):
        markdown, summary = convert_requirement_file_to_markdown("需求.md", "# 登录需求".encode("utf-8"))

        self.assertEqual(markdown, "# 登录需求\n")
        self.assertEqual(summary, "文本文件直接保存为 Markdown 转换稿。")

    def test_text_file_is_wrapped_as_markdown(self):
        markdown, summary = convert_requirement_file_to_markdown("需求.txt", "登录需求".encode("utf-8"))

        self.assertEqual(markdown, "# 需求.txt\n\n登录需求\n")
        self.assertEqual(summary, "文本文件直接保存为 Markdown 转换稿。")

    def test_docx_extracts_headings_lists_and_tables(self):
        raw_bytes = make_docx_bytes()

        markdown, summary = convert_requirement_file_to_markdown("登录需求.docx", raw_bytes)

        self.assertIn("# 登录需求.docx", markdown)
        self.assertIn("# 一级标题", markdown)
        self.assertIn("## 二级标题", markdown)
        self.assertIn("- 支持账号登录", markdown)
        self.assertIn("1. 输入正确密码", markdown)
        self.assertIn("| 字段 | 说明 |", markdown)
        self.assertIn("| --- | --- |", markdown)
        self.assertIn("| username | 登录账号 |", markdown)
        self.assertEqual(summary, "已通过 Python Word 转换器提取正文、标题、列表和表格。")

    def test_docx_preserves_links_code_blocks_and_images(self):
        raw_bytes = make_rich_docx_bytes()
        with tempfile.TemporaryDirectory() as temp_dir:
            assets_dir = Path(temp_dir) / "docmap-1_assets"

            markdown, summary = convert_requirement_file_to_markdown("富文本需求.docx", raw_bytes, assets_dir=assets_dir)

            self.assertIn("[官网](https://example.com)", markdown)
            self.assertIn("```\ncurl https://example.com/api\n```", markdown)
            self.assertIn("![image-1](docmap-1_assets/image-1.png)", markdown)
            self.assertTrue((assets_dir / "image-1.png").exists())
            self.assertEqual(summary, "已通过 Python Word 转换器提取正文、标题、列表和表格、图片 1 个。")

    def test_pdf_merges_pages_without_page_headings_and_removes_noise(self):
        with patch("app.services.requirement_file_converter.fitz.open", return_value=FakePdfDocument()):
            markdown, summary = convert_requirement_file_to_markdown("入园办公人员统计表.pdf", b"%PDF-1.4")

        self.assertIn("# 入园办公人员统计表.pdf", markdown)
        self.assertIn("第一段需求", markdown)
        self.assertIn("第二段需求", markdown)
        self.assertIn("第三段需求", markdown)
        self.assertNotIn("## 第 1 页", markdown)
        self.assertNotIn("公司页眉", markdown)
        self.assertNotIn("保密页脚", markdown)
        self.assertNotIn("\n1\n", markdown)
        self.assertNotIn("Page 2", markdown)
        self.assertEqual(summary, "已通过 PyMuPDF PDF 转换器提取文本，页数 3。")

    def test_pdf_without_text_fails_clearly(self):
        with patch("app.services.requirement_file_converter.fitz.open", return_value=FakePdfDocument(["", "   "])):
            with self.assertRaises(RuntimeError) as caught:
                convert_requirement_file_to_markdown("扫描件.pdf", b"%PDF-1.4")

        self.assertIn("当前不支持 OCR", str(caught.exception))

    def test_pdf_orders_text_blocks_by_position(self):
        document = FakePdfDocument([])
        document.pages = [FakePdfPage("", blocks=[
            (0, 40, 100, 50, "第三行", 3, 0),
            (0, 10, 100, 20, "第一行", 1, 0),
            (0, 25, 100, 35, "第二行", 2, 0),
        ])]
        document.page_count = 1

        with patch("app.services.requirement_file_converter.fitz.open", return_value=document):
            markdown, _summary = convert_requirement_file_to_markdown("排序.pdf", b"%PDF-1.4")

        self.assertLess(markdown.index("第一行"), markdown.index("第二行"))
        self.assertLess(markdown.index("第二行"), markdown.index("第三行"))

    def test_docx_adjacent_list_items_are_not_separated_by_blank_lines(self):
        document = Document()
        document.add_paragraph("支持账号登录", style="List Bullet")
        document.add_paragraph("支持手机登录", style="List Bullet")
        document.add_paragraph("支持微信登录", style="List Bullet")
        buffer = BytesIO()
        document.save(buffer)

        markdown, _ = convert_requirement_file_to_markdown("列表.docx", buffer.getvalue())

        # 三个列表项之间不应有空行（紧凑列表）
        self.assertIn("- 支持账号登录\n- 支持手机登录\n- 支持微信登录", markdown)

    def test_docx_numbered_list_items_are_not_separated_by_blank_lines(self):
        document = Document()
        document.add_paragraph("输入账号", style="List Number")
        document.add_paragraph("输入密码", style="List Number")
        document.add_paragraph("点击登录", style="List Number")
        buffer = BytesIO()
        document.save(buffer)

        markdown, _ = convert_requirement_file_to_markdown("编号.docx", buffer.getvalue())

        self.assertIn("1. 输入账号\n1. 输入密码\n1. 点击登录", markdown)

    def test_docx_list_and_paragraph_spacing_is_correct(self):
        document = Document()
        document.add_paragraph("前言段落")
        document.add_paragraph("功能一", style="List Bullet")
        document.add_paragraph("功能二", style="List Bullet")
        document.add_paragraph("功能三", style="List Bullet")
        document.add_paragraph("结语段落")
        buffer = BytesIO()
        document.save(buffer)

        markdown, _ = convert_requirement_file_to_markdown("混合.docx", buffer.getvalue())

        # 列表内部紧凑
        self.assertIn("- 功能一\n- 功能二\n- 功能三", markdown)
        # 段落与列表之间有空行
        self.assertIn("前言段落\n\n- 功能一", markdown)
        self.assertIn("- 功能三\n\n结语段落", markdown)

    def test_docx_nested_bullet_list_is_indented(self):
        document = Document()
        document.add_paragraph("父级项", style="List Bullet")
        document.add_paragraph("子级项A", style="List Bullet 2")
        document.add_paragraph("子级项B", style="List Bullet 2")
        document.add_paragraph("另一父级", style="List Bullet")
        buffer = BytesIO()
        document.save(buffer)

        markdown, _ = convert_requirement_file_to_markdown("嵌套.docx", buffer.getvalue())

        self.assertIn("- 父级项", markdown)
        self.assertIn("  - 子级项A", markdown)
        self.assertIn("  - 子级项B", markdown)
        self.assertIn("- 另一父级", markdown)

    def test_docx_nested_numbered_list_is_indented(self):
        document = Document()
        document.add_paragraph("步骤一", style="List Number")
        document.add_paragraph("步骤一子步骤", style="List Number 2")
        document.add_paragraph("步骤二", style="List Number")
        buffer = BytesIO()
        document.save(buffer)

        markdown, _ = convert_requirement_file_to_markdown("嵌套编号.docx", buffer.getvalue())

        self.assertIn("1. 步骤一", markdown)
        self.assertIn("  1. 步骤一子步骤", markdown)
        self.assertIn("1. 步骤二", markdown)


def make_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("一级标题", level=1)
    document.add_heading("二级标题", level=2)
    document.add_paragraph("普通段落")
    document.add_paragraph("支持账号登录", style="List Bullet")
    document.add_paragraph("输入正确密码", style="List Number")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "username"
    table.cell(1, 1).text = "登录账号"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_rich_docx_bytes() -> bytes:
    document = Document()
    add_hyperlink_paragraph(document, "访问", "官网", "https://example.com")
    document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("curl https://example.com/api", style="Code")
    with tempfile.TemporaryDirectory() as temp_dir:
        image_path = Path(temp_dir) / "pixel.png"
        image_path.write_bytes(base64.b64decode(ONE_PIXEL_PNG_BASE64))
        document.add_picture(str(image_path))
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()


def add_hyperlink_paragraph(document: Document, prefix: str, text: str, url: str) -> None:
    paragraph = document.add_paragraph(prefix)
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


ONE_PIXEL_PNG_BASE64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="


class FakePdfDocument:
    def __init__(self, page_texts: list[str] | None = None):
        self.pages = [FakePdfPage(text) for text in (page_texts or default_pdf_pages())]
        self.page_count = len(self.pages)

    def __iter__(self):
        return iter(self.pages)


class FakePdfPage:
    def __init__(self, text: str, blocks=None):
        self.text = text
        self.blocks = blocks

    def get_text(self, mode: str = "text"):
        if mode == "blocks":
            if self.blocks is not None:
                return self.blocks
            return [
                (0, index * 20, 100, index * 20 + 10, line, index, 0)
                for index, line in enumerate(self.text.splitlines())
                if line.strip()
            ]
        return self.text


def default_pdf_pages() -> list[str]:
    return [
        "公司页眉\n\n第一段需求\n\n1\n保密页脚",
        "公司页眉\n第二段需求\nPage 2\n保密页脚",
        "公司页眉\n第三段需求\n3 / 3\n保密页脚",
    ]
