from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from docx import Document
from pypdf import PdfReader


@dataclass(frozen=True)
class ParsedRequirementDocument:
    markdown: str
    summary: str


class RequirementFileParserAgent:
    supported_extensions = (".pdf", ".docx", ".doc", ".txt", ".md", ".markdown")

    def parse(self, filename: str, raw_bytes: bytes) -> ParsedRequirementDocument:
        lowered = filename.lower()
        if lowered.endswith(".pdf"):
            return self._parse_pdf(filename, raw_bytes)
        if lowered.endswith((".docx", ".doc")):
            return self._parse_word(filename, raw_bytes)
        if lowered.endswith((".txt", ".md", ".markdown")):
            return self._parse_text(filename, raw_bytes)
        return self._parse_text_preview(filename, raw_bytes)

    def _parse_pdf(self, filename: str, raw_bytes: bytes) -> ParsedRequirementDocument:
        reader = PdfReader(BytesIO(raw_bytes))
        pages: list[str] = []
        for page_index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"## 第 {page_index} 页\n\n{text.strip()}")
        markdown = f"# {filename}\n\n" + "\n\n".join(pages)
        return ParsedRequirementDocument(markdown=markdown.strip() + "\n", summary=f"已解析 PDF，共 {len(reader.pages)} 页。")

    def _parse_word(self, filename: str, raw_bytes: bytes) -> ParsedRequirementDocument:
        document = Document(BytesIO(raw_bytes))
        blocks: list[str] = [f"# {filename}"]
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = (paragraph.style.name or "").lower()
            if "heading 1" in style_name or "标题 1" in style_name:
                blocks.append(f"\n# {text}")
            elif "heading 2" in style_name or "标题 2" in style_name:
                blocks.append(f"\n## {text}")
            elif "heading 3" in style_name or "标题 3" in style_name:
                blocks.append(f"\n### {text}")
            else:
                blocks.append(text)

        for table_index, table in enumerate(document.tables, start=1):
            rows = [[cell.text.strip().replace("\n", "<br>") for cell in row.cells] for row in table.rows]
            if rows:
                blocks.append(f"\n## 表格 {table_index}")
                blocks.extend(_markdown_table(rows))

        markdown = "\n\n".join(blocks)
        return ParsedRequirementDocument(markdown=markdown.strip() + "\n", summary=f"已解析 Word，段落 {len(document.paragraphs)} 个，表格 {len(document.tables)} 个。")

    def _parse_text(self, filename: str, raw_bytes: bytes) -> ParsedRequirementDocument:
        text = _decode_text(raw_bytes)
        if filename.lower().endswith((".md", ".markdown")):
            markdown = text
        else:
            markdown = f"# {filename}\n\n{text.strip()}\n"
        return ParsedRequirementDocument(markdown=markdown.strip() + "\n", summary="已解析文本需求文件。")

    def _parse_text_preview(self, filename: str, raw_bytes: bytes) -> ParsedRequirementDocument:
        text = _decode_text(raw_bytes)
        markdown = f"# {filename}\n\n暂不支持该文件类型的完整结构解析，已保存文本预览。\n\n```\n{text[:4000]}\n```\n"
        return ParsedRequirementDocument(markdown=markdown, summary="文件类型暂不支持完整解析，已生成文本预览。")


requirement_file_parser_agent = RequirementFileParserAgent()


def _decode_text(raw_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return raw_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw_bytes.decode("utf-8", errors="ignore")


def _markdown_table(rows: list[list[str]]) -> list[str]:
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    body = normalized[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return lines
