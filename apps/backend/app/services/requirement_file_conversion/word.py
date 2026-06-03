from dataclasses import dataclass
from io import BytesIO
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from app.services.requirement_file_conversion.common import (
    ConvertedRequirementFile,
    clean_inline_text,
    clean_markdown_inline,
)


@dataclass
class DocxConversionContext:
    assets_dir: Path | None = None
    image_count: int = 0


def convert_docx_bytes_to_markdown(
    filename: str,
    raw_bytes: bytes,
    *,
    assets_dir: Path | None = None,
) -> ConvertedRequirementFile:
    try:
        document = Document(BytesIO(raw_bytes))
    except Exception as exc:  # pragma: no cover - parser details vary by python-docx version
        raise RuntimeError(f"无法转换 {filename}：Word 文件无法解析。") from exc

    context = DocxConversionContext(assets_dir=assets_dir)
    blocks = []
    for element in document.element.body:
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            block = _paragraph_to_markdown(element, document, context)
        elif tag == "tbl":
            block = _table_to_markdown(element, document, context)
        else:
            block = ""
        if block:
            blocks.append(block)

    if not blocks:
        raise RuntimeError(f"无法转换 {filename}：Word 文档未提取到可用文本。")

    markdown = _join_markdown_blocks(blocks) + "\n"
    image_summary = f"、图片 {context.image_count} 个" if context.image_count else ""
    return ConvertedRequirementFile(markdown=markdown, summary=f"已通过 Python Word 转换器提取正文、标题、列表和表格{image_summary}。")


def convert_word_file_to_markdown(input_path: str, *, assets_dir: str | None = None) -> tuple[str, str]:
    source_path = Path(input_path)
    converted = convert_docx_bytes_to_markdown(
        source_path.name,
        source_path.read_bytes(),
        assets_dir=Path(assets_dir) if assets_dir else None,
    )
    return converted.markdown, converted.summary


def _paragraph_to_markdown(element, document: Document, context: DocxConversionContext) -> str:
    paragraph = next((item for item in document.paragraphs if item._element is element), None)
    if paragraph is None:
        return ""
    content = _paragraph_content_to_markdown(paragraph, context)
    if not content:
        return ""

    style_name = paragraph.style.name if paragraph.style is not None else ""
    heading_level = _heading_level(style_name)
    if heading_level:
        return f"{'#' * heading_level} {content}"
    if _is_code_block(style_name):
        return f"```\n{content}\n```"
    if _is_numbered_list(style_name):
        indent = "  " * _list_indent_level(style_name)
        return f"{indent}1. {content}"
    if _is_bullet_list(style_name):
        indent = "  " * _list_indent_level(style_name)
        return f"{indent}- {content}"
    return content


def _table_to_markdown(element, document: Document, context: DocxConversionContext) -> str:
    table = next((item for item in document.tables if item._element is element), None)
    if table is None:
        return ""

    rows = [
        [_cell_to_markdown(cell, context).replace("\n", "<br>") for cell in row.cells]
        for row in table.rows
    ]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return ""

    width = max(len(row) for row in rows)
    normalized_rows = [row + [""] * (width - len(row)) for row in rows]
    header = normalized_rows[0]
    separator = ["---"] * width
    body = normalized_rows[1:]
    markdown_rows = [_markdown_table_row(header), _markdown_table_row(separator)]
    markdown_rows.extend(_markdown_table_row(row) for row in body)
    return "\n".join(markdown_rows)


def _heading_level(style_name: str) -> int | None:
    match = re.match(r"Heading\s+([1-6])$", style_name, flags=re.IGNORECASE)
    if match:
        return int(match.group(1))
    match = re.match(r"标题\s*([1-6])$", style_name)
    if match:
        return int(match.group(1))
    return None


def _list_indent_level(style_name: str) -> int:
    """Return 0-based indentation level from style name."""
    match = re.search(r"\s+(\d+)$", style_name)
    if match:
        return max(0, int(match.group(1)) - 1)
    return 0


def _is_bullet_list(style_name: str) -> bool:
    normalized = style_name.lower()
    return "bullet" in normalized or "项目符号" in style_name


def _is_numbered_list(style_name: str) -> bool:
    normalized = style_name.lower()
    return "number" in normalized or "编号" in style_name


def _is_code_block(style_name: str) -> bool:
    normalized = style_name.lower()
    return "code" in normalized or "preformatted" in normalized or "代码" in style_name


def _paragraph_content_to_markdown(paragraph, context: DocxConversionContext) -> str:
    parts = []
    for child in paragraph._element:
        local_name = child.tag.rsplit("}", 1)[-1]
        if local_name == "r":
            parts.append(_run_to_markdown(child, paragraph.part, context))
        elif local_name == "hyperlink":
            parts.append(_hyperlink_to_markdown(child, paragraph.part, context))
    return clean_markdown_inline("".join(parts))


def _hyperlink_to_markdown(element, part, context: DocxConversionContext) -> str:
    text = clean_inline_text("".join(text_node.text or "" for text_node in element.iter() if text_node.tag == qn("w:t")))
    if not text:
        return ""
    relationship_id = element.get(qn("r:id"))
    if not relationship_id:
        return text
    relationship = part.rels.get(relationship_id)
    if relationship is None:
        return text
    return f"[{_escape_markdown_link_text(text)}]({relationship.target_ref})"


def _run_to_markdown(element, part, context: DocxConversionContext) -> str:
    pieces = []
    text = "".join(text_node.text or "" for text_node in element.iter() if text_node.tag == qn("w:t"))
    if text:
        pieces.append(text)
    for blip in element.iter(qn("a:blip")):
        relationship_id = blip.get(qn("r:embed"))
        image_markdown = _image_to_markdown(relationship_id, part, context)
        if image_markdown:
            pieces.append(image_markdown)
    return "".join(pieces)


def _image_to_markdown(relationship_id: str | None, part, context: DocxConversionContext) -> str:
    if not relationship_id or context.assets_dir is None:
        return ""
    image_part = part.related_parts.get(relationship_id)
    if image_part is None:
        return ""

    context.assets_dir.mkdir(parents=True, exist_ok=True)
    context.image_count += 1
    extension = _image_extension(image_part.content_type)
    filename = f"image-{context.image_count}{extension}"
    image_path = context.assets_dir / filename
    image_path.write_bytes(image_part.blob)
    return f"![image-{context.image_count}]({context.assets_dir.name}/{filename})"


def _image_extension(content_type: str) -> str:
    return {
        "image/jpeg": ".jpg",
        "image/png": ".png",
        "image/gif": ".gif",
        "image/bmp": ".bmp",
        "image/tiff": ".tiff",
        "image/svg+xml": ".svg",
    }.get(content_type, ".bin")


def _cell_to_markdown(cell, context: DocxConversionContext) -> str:
    blocks = []
    for paragraph in cell.paragraphs:
        content = _paragraph_content_to_markdown(paragraph, context)
        if content:
            blocks.append(content)
    return "<br>".join(blocks)


def _markdown_table_row(cells: list[str]) -> str:
    escaped = [cell.replace("|", "\\|") for cell in cells]
    return "| " + " | ".join(escaped) + " |"


def _escape_markdown_link_text(text: str) -> str:
    return text.replace("[", "\\[").replace("]", "\\]")


def _list_block_type(block: str) -> str | None:
    if re.match(r"^\s*[-*+]\s", block):
        return "bullet"
    if re.match(r"^\s*\d+\.\s", block):
        return "ordered"
    return None


def _join_markdown_blocks(blocks: list[str]) -> str:
    if not blocks:
        return ""
    groups: list[str] = []
    i = 0
    while i < len(blocks):
        block = blocks[i]
        block_type = _list_block_type(block)
        if block_type is not None:
            list_items = [block]
            while i + 1 < len(blocks) and _list_block_type(blocks[i + 1]) == block_type:
                i += 1
                list_items.append(blocks[i])
            groups.append("\n".join(list_items))
        else:
            groups.append(block)
        i += 1
    return "\n\n".join(groups)
