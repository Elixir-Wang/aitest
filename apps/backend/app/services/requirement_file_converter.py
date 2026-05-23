from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from io import BytesIO
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
import fitz


@dataclass(frozen=True)
class ConvertedRequirementFile:
    markdown: str
    summary: str


@dataclass
class DocxConversionContext:
    assets_dir: Path | None = None
    image_count: int = 0


def convert_requirement_file_to_markdown(
    filename: str,
    raw_bytes: bytes,
    *,
    assets_dir: Path | None = None,
) -> tuple[str, str]:
    lowered = filename.lower()
    if lowered.endswith((".md", ".markdown", ".txt")):
        return _normalize_text_markdown(filename, raw_bytes), "文本文件直接保存为 Markdown 转换稿。"
    if lowered.endswith(".pdf"):
        converted = _convert_pdf_to_markdown(filename, raw_bytes)
        return converted.markdown, converted.summary
    if lowered.endswith((".doc", ".docx")):
        converted = _convert_docx_to_markdown(filename, raw_bytes, assets_dir=assets_dir)
        return converted.markdown, converted.summary
    preview = _decode_text(raw_bytes)[:4000]
    markdown = f"# {filename}\n\n暂不支持该文件类型的完整结构解析，已保存文本预览。\n\n```\n{preview}\n```\n"
    return markdown, "文件类型暂不支持完整解析，已生成文本预览。"


def _normalize_text_markdown(filename: str, raw_bytes: bytes) -> str:
    text = _decode_text(raw_bytes).strip()
    if filename.lower().endswith((".md", ".markdown")):
        return text + "\n"
    return f"# {filename}\n\n{text}\n"


def _convert_pdf_to_markdown(filename: str, raw_bytes: bytes) -> ConvertedRequirementFile:
    try:
        document = fitz.open(stream=raw_bytes, filetype="pdf")
    except Exception as exc:  # pragma: no cover - parser details vary by PyMuPDF version
        raise RuntimeError(f"无法转换 {filename}：PDF 文件无法解析。") from exc

    pages = []
    for page in document:
        pages.append(_extract_pdf_page_lines(page))

    repeated_noise = _detect_repeated_pdf_noise(pages)
    content_blocks = []
    for lines in pages:
        cleaned = [line for line in lines if line not in repeated_noise and not _is_page_number_line(line)]
        block = "\n".join(cleaned).strip()
        if block:
            content_blocks.append(block)

    if not content_blocks:
        raise RuntimeError(f"无法转换 {filename}：PDF 未提取到可用文本，可能是扫描件，当前不支持 OCR。")

    markdown = f"# {filename}\n\n" + "\n\n".join(content_blocks) + "\n"
    return ConvertedRequirementFile(markdown=markdown, summary=f"已通过 PyMuPDF PDF 转换器提取文本，页数 {document.page_count}。")


def _convert_docx_to_markdown(filename: str, raw_bytes: bytes, *, assets_dir: Path | None = None) -> ConvertedRequirementFile:
    try:
        document = Document(BytesIO(raw_bytes))
    except Exception as exc:  # pragma: no cover - parser details vary by python-docx version
        raise RuntimeError(f"无法转换 {filename}：Word 文件无法解析。") from exc

    context = DocxConversionContext(assets_dir=assets_dir)
    blocks = []
    for element in document.element.body:
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            block = _paragraph_to_markdown(element, document, context)
        elif tag == "tbl":
            block = _table_to_markdown(element, document, context)
        else:
            block = ""
        if block:
            blocks.append(block)

    if not blocks:
        raise RuntimeError(f"无法转换 {filename}：Word 文档未提取到可用文本。")

    markdown = f"# {filename}\n\n" + _join_markdown_blocks(blocks) + "\n"
    image_summary = f"、图片 {context.image_count} 个" if context.image_count else ""
    return ConvertedRequirementFile(markdown=markdown, summary=f"已通过 Python Word 转换器提取正文、标题、列表和表格{image_summary}。")


def _paragraph_to_markdown(element, document: Document, context: DocxConversionContext) -> str:
    paragraph = next((item for item in document.paragraphs if item._element is element), None)
    if paragraph is None:
        return ""
    content = _paragraph_content_to_markdown(paragraph, context)
    if not content:
        return ""

    style_name = paragraph.style.name if paragraph.style is not None else ""
    heading_level = _heading_level(style_name)
    if heading_level:
        return f"{'#' * heading_level} {content}"
    if _is_code_block(style_name):
        return f"```\n{content}\n```"
    if _is_numbered_list(style_name):
        return f"1. {content}"
    if _is_bullet_list(style_name):
        return f"- {content}"
    return content


def _table_to_markdown(element, document: Document, context: DocxConversionContext) -> str:
    table = next((item for item in document.tables if item._element is element), None)
    if table is None:
        return ""

    rows = [
        [_cell_to_markdown(cell, context).replace("\n", "<br>") for cell in row.cells]
        for row in table.rows
    ]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return ""

    width = max(len(row) for row in rows)
    normalized_rows = [row + [""] * (width - len(row)) for row in rows]
    header = normalized_rows[0]
    separator = ["---"] * width
    body = normalized_rows[1:]
    markdown_rows = [_markdown_table_row(header), _markdown_table_row(separator)]
    markdown_rows.extend(_markdown_table_row(row) for row in body)
    return "\n".join(markdown_rows)


def _heading_level(style_name: str) -> int | None:
    match = re.match(r"Heading\s+([1-6])$", style_name, flags=re.IGNORECASE)
    if match:
        return int(match.group(1))
    match = re.match(r"标题\s*([1-6])$", style_name)
    if match:
        return int(match.group(1))
    return None


def _is_bullet_list(style_name: str) -> bool:
    normalized = style_name.lower()
    return "bullet" in normalized or "项目符号" in style_name


def _is_numbered_list(style_name: str) -> bool:
    normalized = style_name.lower()
    return "number" in normalized or "编号" in style_name


def _is_code_block(style_name: str) -> bool:
    normalized = style_name.lower()
    return "code" in normalized or "preformatted" in normalized or "代码" in style_name


def _paragraph_content_to_markdown(paragraph, context: DocxConversionContext) -> str:
    parts = []
    for child in paragraph._element:
        local_name = child.tag.rsplit("}", 1)[-1]
        if local_name == "r":
            parts.append(_run_to_markdown(child, paragraph.part, context))
        elif local_name == "hyperlink":
            parts.append(_hyperlink_to_markdown(child, paragraph.part, context))
    return _clean_markdown_inline("".join(parts))


def _hyperlink_to_markdown(element, part, context: DocxConversionContext) -> str:
    text = _clean_inline_text("".join(text_node.text or "" for text_node in element.iter() if text_node.tag == qn("w:t")))
    if not text:
        return ""
    relationship_id = element.get(qn("r:id"))
    if not relationship_id:
        return text
    relationship = part.rels.get(relationship_id)
    if relationship is None:
        return text
    return f"[{_escape_markdown_link_text(text)}]({relationship.target_ref})"


def _run_to_markdown(element, part, context: DocxConversionContext) -> str:
    pieces = []
    text = "".join(text_node.text or "" for text_node in element.iter() if text_node.tag == qn("w:t"))
    if text:
        pieces.append(text)
    for blip in element.iter(qn("a:blip")):
        relationship_id = blip.get(qn("r:embed"))
        image_markdown = _image_to_markdown(relationship_id, part, context)
        if image_markdown:
            pieces.append(image_markdown)
    return "".join(pieces)


def _image_to_markdown(relationship_id: str | None, part, context: DocxConversionContext) -> str:
    if not relationship_id or context.assets_dir is None:
        return ""
    image_part = part.related_parts.get(relationship_id)
    if image_part is None:
        return ""

    context.assets_dir.mkdir(parents=True, exist_ok=True)
    context.image_count += 1
    extension = _image_extension(image_part.content_type)
    filename = f"image-{context.image_count}{extension}"
    image_path = context.assets_dir / filename
    image_path.write_bytes(image_part.blob)
    return f"![image-{context.image_count}]({context.assets_dir.name}/{filename})"


def _image_extension(content_type: str) -> str:
    return {
        "image/jpeg": ".jpg",
        "image/png": ".png",
        "image/gif": ".gif",
        "image/bmp": ".bmp",
        "image/tiff": ".tiff",
        "image/svg+xml": ".svg",
    }.get(content_type, ".bin")


def _cell_to_markdown(cell, context: DocxConversionContext) -> str:
    blocks = []
    for paragraph in cell.paragraphs:
        content = _paragraph_content_to_markdown(paragraph, context)
        if content:
            blocks.append(content)
    return "<br>".join(blocks)


def _markdown_table_row(cells: list[str]) -> str:
    escaped = [cell.replace("|", "\\|") for cell in cells]
    return "| " + " | ".join(escaped) + " |"


def _extract_pdf_page_lines(page) -> list[str]:
    blocks = page.get_text("blocks") or []
    text_blocks = []
    for block in blocks:
        if len(block) < 5:
            continue
        x0, y0, _x1, _y1, text = block[:5]
        if not isinstance(text, str) or not text.strip():
            continue
        text_blocks.append((round(float(y0), 1), round(float(x0), 1), text))

    if text_blocks:
        ordered_text = "\n\n".join(text for _y, _x, text in sorted(text_blocks))
        return _clean_pdf_page_lines(ordered_text.splitlines())

    raw_text = page.get_text("text") or ""
    return _clean_pdf_page_lines(raw_text.splitlines())


def _clean_pdf_page_lines(lines: list[str]) -> list[str]:
    cleaned = []
    previous_blank = False
    for raw_line in lines:
        line = _clean_inline_text(raw_line)
        if not line:
            previous_blank = True
            continue
        if previous_blank and cleaned and cleaned[-1] != "":
            cleaned.append("")
        cleaned.append(line)
        previous_blank = False
    while "" in cleaned[:1]:
        cleaned.pop(0)
    while "" in cleaned[-1:]:
        cleaned.pop()
    return cleaned


def _detect_repeated_pdf_noise(pages: list[list[str]]) -> set[str]:
    if len(pages) < 3:
        return set()
    candidates = []
    for lines in pages:
        meaningful = [line for line in lines if line and not _is_page_number_line(line)]
        candidates.extend(meaningful[:1])
        candidates.extend(meaningful[-1:])
    threshold = max(2, len(pages) // 2 + 1)
    return {line for line, count in Counter(candidates).items() if count >= threshold}


def _is_page_number_line(line: str) -> bool:
    normalized = line.strip()
    return bool(
        re.fullmatch(r"\d+", normalized)
        or re.fullmatch(r"-\s*\d+\s*-", normalized)
        or re.fullmatch(r"第\s*\d+\s*页", normalized)
        or re.fullmatch(r"page\s+\d+", normalized, flags=re.IGNORECASE)
        or re.fullmatch(r"\d+\s*/\s*\d+", normalized)
    )


def _clean_inline_text(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip()


def _clean_markdown_inline(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text.replace("\u00a0", " ")).strip()


def _escape_markdown_link_text(text: str) -> str:
    return text.replace("[", "\\[").replace("]", "\\]")


def _list_block_type(block: str) -> str | None:
    """Return the list type of a markdown block, or None if not a list item."""
    if re.match(r"^\s*[-*+]\s", block):
        return "bullet"
    if re.match(r"^\s*\d+\.\s", block):
        return "ordered"
    return None


def _join_markdown_blocks(blocks: list[str]) -> str:
    """Join blocks with double newlines, but keep consecutive same-type list items tight (single newline)."""
    if not blocks:
        return ""
    groups: list[str] = []
    i = 0
    while i < len(blocks):
        block = blocks[i]
        block_type = _list_block_type(block)
        if block_type is not None:
            list_items = [block]
            while i + 1 < len(blocks) and _list_block_type(blocks[i + 1]) == block_type:
                i += 1
                list_items.append(blocks[i])
            groups.append("\n".join(list_items))
        else:
            groups.append(block)
        i += 1
    return "\n\n".join(groups)


def _decode_text(raw_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return raw_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw_bytes.decode("utf-8", errors="ignore")
